"""The midseason draft runs in reverse order of the standings after Week 7.

scripts/export_current.py publishes the order with the upcoming draft; the
Drafts page and the midseason draft board export both read it from there.
"""

import base64
import json
from io import BytesIO

from openpyxl import load_workbook

from scripts.export_current import midseason_draft_order
from tests.test_api import FakeRepo, transaction

TEAMS = [
    {'abbrev': 'GSA', 'name': 'Team GSA', 'owner': 'A'},
    {'abbrev': 'WJK', 'name': 'Team WJK', 'owner': 'B'},
]


def _web_dir(tmp_path):
    web_dir = tmp_path / 'web'
    week_dir = web_dir / 'data' / 'seasons' / '2026' / 'weeks'
    week_dir.mkdir(parents=True)
    (week_dir / 'week_1.json').write_text(
        json.dumps(
            {
                'week': 1,
                'has_scores': True,
                'teams': [
                    {'abbrev': 'GSA', 'name': 'Team GSA', 'total_score': 120},
                    {'abbrev': 'WJK', 'name': 'Team WJK', 'total_score': 90},
                ],
                'matchups': [
                    {
                        'team1': {'abbrev': 'GSA', 'total_score': 120},
                        'team2': {'abbrev': 'WJK', 'total_score': 90},
                    }
                ],
            }
        )
    )
    return web_dir


def _data(completed_through):
    return {
        'teams': TEAMS,
        'current_week': completed_through + 1,
        'standings': [],
        'hall_of_fame': {'completed_through': {'2026': completed_through}},
    }


def test_order_is_provisional_until_week_7_is_final(tmp_path):
    order = midseason_draft_order(_data(1), _web_dir(tmp_path), 2026)

    assert order == {'order': ['WJK', 'GSA'], 'through_week': 1, 'final': False}


def test_order_stops_at_week_7(tmp_path):
    order = midseason_draft_order(_data(9), _web_dir(tmp_path), 2026)

    assert order['through_week'] == 7
    assert order['final'] is True


def test_no_order_before_any_week_is_final(tmp_path):
    assert midseason_draft_order(_data(0), _web_dir(tmp_path), 2026) is None


ORDER = ['WJK', 'GSA']


def _picks():
    picks = []
    for draft_type, rounds in (('waiver', 6), ('waiver_taxi', 4)):
        for round_number in range(1, rounds + 1):
            for team in ORDER:
                picks.append(
                    {
                        'year': '2026',
                        'round': round_number,
                        'draft_type': draft_type,
                        'original_team': team,
                        'current_owner': team,
                        'previous_owners': [],
                    }
                )
    picks[0].update(current_owner='GSA', previous_owners=['WJK'])
    return {'picks': picks}


def _live(final):
    return {
        'upcoming_drafts': [
            {
                'name': '2026 Midseason Draft',
                'type': 'waiver',
                'order_basis': {'through_week': 7 if final else 6, 'final': final},
                'rounds': [
                    {
                        'round': 1,
                        'picks': [
                            {'original_team': 'GSA', 'pick_number': '1.02'},
                            {'original_team': 'WJK', 'pick_number': '1.01'},
                        ],
                    }
                ],
            }
        ]
    }


def _export(monkeypatch, live):
    monkeypatch.setenv('TEAM_PASSWORD_CGK', 'pw')
    FakeRepo(
        {
            'data/teams.json': {'teams': TEAMS},
            'data/draft_picks.json': _picks(),
            transaction.SITE_LIVE_PATH: live,
        }
    ).install(monkeypatch)
    return transaction.handle_export_workbook(
        {'team': 'CGK', 'password': 'pw', 'export': 'download_midseason_draft_board'}
    )


def test_midseason_board_follows_the_published_order_with_trades(monkeypatch):
    status, body = _export(monkeypatch, _live(final=True))

    assert status == 200, body
    assert body['filename'] == '2026_Midseason_Draft_Board.xlsx'
    workbook = load_workbook(BytesIO(base64.b64decode(body['content_base64'])))
    board = workbook['2026 Midseason Draft']
    assert board['A1'].value == '2026 Midseason Draft'
    assert board['F1'].value == 'Order: reverse standings through Week 7 (final)'
    assert [board['B4'].value, board['B5'].value] == ['GSA (via WJK)', 'GSA']
    ledger = list(workbook['Pick Ledger'].iter_rows(min_row=2, values_only=True))
    assert {row[0] for row in ledger} == {'Midseason', 'Taxi'}
    assert len(ledger) == 20
    workbook.close()


def test_midseason_board_marks_a_provisional_order(monkeypatch):
    status, body = _export(monkeypatch, _live(final=False))

    assert status == 200, body
    workbook = load_workbook(BytesIO(base64.b64decode(body['content_base64'])))
    note = workbook['2026 Midseason Draft']['F1'].value
    assert note == 'Order: reverse standings through Week 6 (provisional until Week 7 is final)'
    workbook.close()


def test_midseason_board_needs_a_published_draft(monkeypatch):
    status, body = _export(monkeypatch, {'upcoming_drafts': []})

    assert status == 500
    assert 'No upcoming 2026 midseason draft' in body['error']
